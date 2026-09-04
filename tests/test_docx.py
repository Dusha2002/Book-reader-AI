from pathlib import Path

from docx import Document

from bookai.parsers.docx import load_docx, save_docx


def test_docx_roundtrip(tmp_path: Path):
    src = tmp_path / "book.docx"
    d = Document()
    d.add_heading("Chapter One", level=1)
    d.add_paragraph("Hello world.")
    d.save(src)
    doc = load_docx(src)
    assert doc.segments[-1].chapter == "Chapter One"
    out = tmp_path / "ru.docx"
    save_docx(doc, {doc.segments[-1].id: "Привет, мир."}, out)
    d2 = Document(out)
    assert "Привет, мир." in [p.text for p in d2.paragraphs]
