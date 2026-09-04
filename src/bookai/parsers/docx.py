from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document

from ..models import BookDocument, Segment


@dataclass
class DocxPayload:
    segment_index: dict[str, int]


def load_docx(path: Path) -> BookDocument:
    doc = Document(str(path))
    segments: list[Segment] = []
    segment_index: dict[str, int] = {}
    chapter = path.stem
    idx = 0
    for pos, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        if "heading" in style or "заголов" in style:
            chapter = text[:200]
        sid = f"s{idx:06d}"
        segments.append(Segment(sid, text, f"paragraph:{pos}", chapter=chapter))
        segment_index[sid] = pos
        idx += 1
    return BookDocument(path, "docx", segments, DocxPayload(segment_index))


def _replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def save_docx(document: BookDocument, translations: dict[str, str], output: Path) -> None:
    payload = document.payload
    assert isinstance(payload, DocxPayload)
    doc = Document(str(document.source))
    for sid, translated in translations.items():
        pos = payload.segment_index.get(sid)
        if pos is not None and pos < len(doc.paragraphs):
            _replace_paragraph_text(doc.paragraphs[pos], translated)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
